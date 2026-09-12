"""Geometric QA for the pitch deck.

Checks three things the eye would otherwise catch: shapes outside the slide,
text boxes whose content is taller than the box, and overlapping text boxes.
Uses Liberation Sans, which is metric-compatible with Arial.
"""

import sys
from pptx import Presentation
from pptx.util import Emu
from PIL import ImageFont

EMU_IN = 914400
W_IN, H_IN = 13.333, 7.5
REG = "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"
BOLD = "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"
PX_PER_PT = 4  # supersample for metric accuracy

_cache = {}


def font(size_pt, bold):
    key = (round(size_pt, 1), bold)
    if key not in _cache:
        _cache[key] = ImageFont.truetype(BOLD if bold else REG,
                                         int(size_pt * PX_PER_PT))
    return _cache[key]


def text_width_pt(text, size_pt, bold):
    return font(size_pt, bold).getlength(text) / PX_PER_PT


def wrapped_lines(text, size_pt, bold, width_pt):
    """Greedy wrap, mirroring how a renderer breaks on spaces."""
    total = 0
    for hard in text.split("\n"):
        words, line = hard.split(), ""
        n = 0
        for w in words:
            trial = f"{line} {w}".strip()
            if text_width_pt(trial, size_pt, bold) <= width_pt or not line:
                line = trial
            else:
                n += 1
                line = w
        total += n + 1
    return total


def para_metrics(p):
    runs = p.runs
    if not runs:
        return None
    size = runs[0].font.size.pt if runs[0].font.size else 18.0
    bold = bool(runs[0].font.bold)
    text = "".join(r.text for r in runs)
    ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.0
    sb = p.space_before.pt if p.space_before else 0.0
    sa = p.space_after.pt if p.space_after else 0.0
    return text, size, bold, ls, sb, sa


problems = []
prs = Presentation(sys.argv[1])

for si, slide in enumerate(prs.slides, 1):
    boxes = []
    for shape in slide.shapes:
        L = shape.left / EMU_IN
        T = shape.top / EMU_IN
        Wd = shape.width / EMU_IN
        Ht = shape.height / EMU_IN

        if L < -0.01 or T < -0.01 or L + Wd > W_IN + 0.01 or T + Ht > H_IN + 0.01:
            problems.append(
                f"slide {si}: OUT OF BOUNDS  "
                f"{shape.shape_type}  x={L:.2f} y={T:.2f} w={Wd:.2f} h={Ht:.2f}")

        if not shape.has_text_frame or not shape.text_frame.text.strip():
            continue

        tf = shape.text_frame
        avail_w_pt = Wd * 72
        needed_pt = 0.0
        for p in tf.paragraphs:
            m = para_metrics(p)
            if not m:
                continue
            text, size, bold, ls, sb, sa = m
            nlines = wrapped_lines(text, size, bold, avail_w_pt)
            needed_pt += sb + sa + nlines * size * 1.2 * ls
        needed_in = needed_pt / 72
        if needed_in > Ht + 0.02:
            problems.append(
                f"slide {si}: TEXT OVERFLOW  needs {needed_in:.2f}in in "
                f"{Ht:.2f}in box (y={T:.2f})  :: "
                f"{tf.text[:58].replace(chr(10), ' ')!r}")
        # record actual painted extent for overlap testing
        boxes.append((L, T, Wd, max(Ht, needed_in),
                      tf.text[:40].replace("\n", " ")))

    for i in range(len(boxes)):
        for j in range(i + 1, len(boxes)):
            ax, ay, aw, ah, at = boxes[i]
            bx, by, bw, bh, bt = boxes[j]
            ox = min(ax + aw, bx + bw) - max(ax, bx)
            oy = min(ay + ah, by + bh) - max(ay, by)
            if ox > 0.05 and oy > 0.05:
                problems.append(
                    f"slide {si}: OVERLAP {ox:.2f}x{oy:.2f}in  "
                    f"{at!r} vs {bt!r}")

print(f"{len(prs.slides._sldIdLst)} slides checked")
if problems:
    print(f"\n{len(problems)} PROBLEM(S):\n")
    for p in problems:
        print(" ", p)
else:
    print("\nNo bounds, overflow or overlap problems found.")
