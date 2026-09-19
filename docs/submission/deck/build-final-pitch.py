"""Build the editable seven-slide finals deck from the official URIIS template.

The seven core template slides are retained as the base. Their header, URIIS logo,
footer and master remain intact. Instructional content is removed and replaced with
native PowerPoint shapes and Arial text for Canva compatibility.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches as In, Pt


ROOT = Path(__file__).resolve().parents[3]
DECK = Path(__file__).resolve().parent
ASSETS = DECK / "assets"
TEMPLATE = ROOT / "docs/source/final-pitch-deck-sample.pptx"
OUTPUT = DECK / "mytakziah-final-pitch.pptx"

FONT = "Arial"
INK = RGBColor(0x28, 0x19, 0x13)
TERRA = RGBColor(0xBC, 0x4D, 0x1A)
OCHRE = RGBColor(0xC7, 0x89, 0x2D)
SLATE = RGBColor(0x37, 0x6D, 0x7C)
MUTED = RGBColor(0x6E, 0x63, 0x5C)
HAIR = RGBColor(0xE2, 0xDB, 0xD4)
SOFT = RGBColor(0xF9, 0xF3, 0xEA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

HEADERS = [
    ("OPEN STRONG", "1. The Idea"),
    ("WHY THIS MATTERS", "2. The Problem"),
    ("MAKE IT EASY TO UNDERSTAND", "3. Our Solution"),
    ("DIFFERENTIATION", "4. What Makes It Innovative?"),
    ("CAN THIS REALLY BE DONE?", "5. How It Works & Feasibility"),
    ("SHOW THE VALUE", "6. Impact & Potential"),
    ("FINISH MEMORABLY", "7. Why Us + Closing Vision"),
]


def remove_slide(prs, index):
    slide_id = prs.slides._sldIdLst[index]
    prs.part.drop_rel(slide_id.rId)
    del prs.slides._sldIdLst[index]


def remove_shape(shape):
    shape._element.getparent().remove(shape._element)


def set_first_run(shape, value):
    paragraph = shape.text_frame.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = value
    else:
        paragraph.text = value


def prepare_template(prs):
    for index in reversed(range(len(prs.slides))):
        if index not in range(3, 10):
            remove_slide(prs, index)

    for number, slide in enumerate(prs.slides, 1):
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = WHITE
        for shape in list(slide.shapes):
            top = shape.top / 914400
            keep = top < 1.60 or top > 6.85
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and top < 1.20:
                keep = True
            if not keep:
                remove_shape(shape)

        for shape in slide.shapes:
            if (
                shape.has_text_frame
                and shape.left == 0
                and shape.top == 0
                and shape.width == prs.slide_width
                and shape.height == prs.slide_height
            ):
                shape.text_frame.clear()

        top_text = sorted(
            [
                shape
                for shape in slide.shapes
                if shape.has_text_frame
                and shape.top / 914400 < 1.60
                and shape.height / 914400 < 1.00
            ],
            key=lambda shape: shape.top,
        )
        if len(top_text) >= 2:
            set_first_run(top_text[0], HEADERS[number - 1][0])
            set_first_run(top_text[1], HEADERS[number - 1][1])

        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            top = shape.top / 914400
            if top > 6.85 and text.isdigit():
                set_first_run(shape, str(number))
            if top > 6.85:
                shape.top = In(7.05)
                shape.height = In(0.30)


def text_box(
    slide,
    text,
    x,
    y,
    w,
    h,
    size=18,
    color=INK,
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin=0.02,
):
    shape = slide.shapes.add_textbox(In(x), In(y), In(w), In(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = In(margin)
    frame.margin_right = In(margin)
    frame.margin_top = In(margin)
    frame.margin_bottom = In(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    paragraph.line_spacing = 1.04
    run = paragraph.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def rect(slide, x, y, w, h, fill=WHITE, line=HAIR, radius=True):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, In(x), In(y), In(w), In(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(0.8)
    return shape


def rule(slide, x, y, w, color=HAIR, width=1.0):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, In(x), In(y), In(x + w), In(y)
    )
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    return shape


def dot(slide, x, y, diameter=0.26, color=SLATE, width=3.5):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, In(x), In(y), In(diameter), In(diameter)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    return shape


def picture(slide, name, x, y, w=None, h=None):
    kwargs = {}
    if w is not None:
        kwargs["width"] = In(w)
    if h is not None:
        kwargs["height"] = In(h)
    return slide.shapes.add_picture(str(ASSETS / name), In(x), In(y), **kwargs)


def source(slide, value):
    text_box(slide, value, 0.75, 6.83, 11.55, 0.22, 9.0, MUTED)


def notes(slide, script, evidence):
    slide.notes_slide.notes_text_frame.text = (
        f"PRESENTER SCRIPT\n\n{script}\n\nEVIDENCE AND LIMITS\n\n{evidence}"
    )


prs = Presentation(str(TEMPLATE))
prepare_template(prs)

# Slide 1: The Idea
s = prs.slides[0]
picture(s, "mark.png", 0.78, 1.92, h=0.76)
text_box(s, "MyTakziah", 1.78, 1.80, 6.0, 0.92, 54, INK, True)
text_box(s, "One Family Record. Clearer Next Steps.", 0.78, 2.92, 7.25, 1.04, 29, TERRA, True)
text_box(
    s,
    "A proposed assistant helping bereaved families prepare for JPN, KWSP and PERKESO.",
    0.78,
    4.08,
    7.0,
    0.80,
    18,
    MUTED,
)
rule(s, 0.78, 5.42, 6.95)
text_box(
    s,
    "Team Survive | Universiti Malaya and Asia Pacific University",
    0.78,
    5.58,
    7.2,
    0.32,
    16,
    INK,
    True,
)
text_box(
    s,
    "Anderson Ling | Hee Zi Jie | Wong Xuan Rui | Yow Li Hao",
    0.78,
    5.98,
    7.0,
    0.34,
    12,
    MUTED,
)
rect(s, 8.44, 2.58, 1.78, 1.52, SOFT)
picture(s, "mark.png", 8.98, 2.83, h=0.55)
text_box(s, "Family Record", 8.60, 3.48, 1.46, 0.28, 13, INK, True, PP_ALIGN.CENTER)
for y, label in [(2.38, "JPN"), (3.41, "KWSP"), (4.44, "PERKESO")]:
    connector = s.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, In(10.22), In(3.34), In(11.05), In(y + 0.25)
    )
    connector.line.color.rgb = SLATE
    connector.line.width = Pt(1.4)
    rect(s, 11.05, y, 1.18, 0.50, WHITE)
    text_box(s, label, 11.05, y + 0.07, 1.18, 0.28, 14, SLATE, True, PP_ALIGN.CENTER)
text_box(
    s,
    "Student Proposal\nNot an Official Government Service",
    8.55,
    5.70,
    3.70,
    0.55,
    10.5,
    SLATE,
    True,
    PP_ALIGN.RIGHT,
)
notes(
    s,
    "Good afternoon. We are Team Survive, from Universiti Malaya and Asia Pacific University. MyTakziah is our proposed assistant for bereaved families: one family record, clearer next steps across death registration, EPF savings and SOCSO benefits.",
    "Team identity and selected scope are recorded in docs/decisions.md. This is a proposed service, not a live government integration.",
)

# Slide 2: The Problem
s = prs.slides[1]
text_box(s, "After a Death, the Paperwork Still Repeats", 0.78, 1.68, 11.4, 0.54, 30, INK, True)
rule(s, 0.78, 2.52, 3.55, TERRA, 3.4)
text_box(s, "198,992", 0.78, 2.70, 3.2, 1.02, 57, TERRA, True)
text_box(s, "deaths recorded in Malaysia in 2024", 0.78, 3.82, 3.25, 0.70, 17, INK)
for y, agency in [(2.62, "JPN"), (3.47, "KWSP"), (4.32, "PERKESO")]:
    rule(s, 4.45, y, 7.78)
    text_box(s, agency, 4.45, y + 0.22, 1.25, 0.35, 18, SLATE, True)
    rect(s, 5.92, y + 0.12, 2.64, 0.56, SOFT)
    text_box(s, "Identity", 6.10, y + 0.25, 2.25, 0.28, 14, INK, True)
    rect(s, 8.80, y + 0.12, 3.42, 0.56, SOFT)
    text_box(s, "Proof of Death", 8.98, y + 0.25, 3.05, 0.28, 14, INK, True)
rect(s, 0.78, 5.45, 11.45, 1.04, WHITE, RGBColor(0xDC, 0xA7, 0x89))
text_box(s, "A Deadline Families Cannot Afford to Miss", 1.00, 5.58, 4.10, 0.50, 14, INK, True)
text_box(
    s,
    "EPF Death Assistance: application within six months, death before 60, and other conditions. Not a deadline for withdrawing all EPF savings.",
    5.25,
    5.60,
    6.63,
    0.60,
    13.0,
    MUTED,
)
source(s, "Sources: DOSM, Vital Statistics, Malaysia, 2025; DOSM, Current Population Estimates, 2026; KWSP, Death Withdrawal.")
notes(
    s,
    "Malaysia recorded almost two hundred thousand deaths in 2024. Its population is also ageing: the share aged sixty-five and over reached eight point four per cent in 2026. Our proposed first user is a surviving spouse handling the paperwork, with a family member helping if needed. The forms show the burden. Identity and death information recur across agencies; relationship and payment details recur across the benefit claims. Timing matters. EPF's discretionary two-thousand-five-hundred-ringgit Death Assistance requires an application within six months and the member's death before sixty, among other conditions. That is one concrete deadline. It is not a deadline for withdrawing all EPF savings.",
    "198,992 is deaths in Malaysia in 2024, not a customer count. The surviving-spouse segment is an assumption. The six-month condition applies to discretionary Death Assistance, not all EPF savings.",
)

# Slide 3: Our Solution
s = prs.slides[2]
text_box(s, "Provide Shared Details Once. Prepare the Next Step.", 0.78, 1.68, 11.6, 0.54, 28, INK, True)
for i, (label, detail) in enumerate(
    [
        ("Reuse", "Shared details in relevant draft forms."),
        ("Review", "Requirements, known deadlines and questions that need checking."),
        ("Submit Officially", "The family uses official channels. Agencies decide."),
    ]
):
    y = 2.55 + i * 0.88
    text_box(s, label, 0.78, y, 2.6, 0.30, 17, TERRA, True)
    text_box(s, detail, 0.78, y + 0.36, 3.0, 0.44, 11.6, MUTED)
    if i < 2:
        rule(s, 0.78, y + 0.78, 3.05)
picture(s, "qr-logo.png", 0.82, 4.78, h=1.60)
text_box(s, "Interactive Prototype", 2.60, 5.19, 1.42, 0.44, 13, INK, True)
text_box(s, "mytakziah.vercel.app", 2.60, 5.72, 1.42, 0.42, 9.8, MUTED)

rect(s, 4.18, 2.53, 3.05, 3.83, SOFT)
pic = picture(s, "live-step2.png", 4.18, 2.53, 3.05, 3.83)
pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = 0.14, 0.30, 0.46, 0.02
rect(s, 7.52, 2.53, 4.72, 2.20, SOFT)
pic = picture(s, "screen-timeline.png", 7.52, 2.53, 4.72, 2.20)
pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = 0.00, 0.02, 0.00, 0.47
text_box(s, "INTERACTIVE PROTOTYPE. AGENCY RESPONSES SIMULATED.", 7.58, 4.94, 4.58, 0.42, 11.5, SLATE, True)
text_box(
    s,
    "Required originals, signatures and thumbprints remain where applicable.",
    7.58,
    5.40,
    4.55,
    0.46,
    12.5,
    MUTED,
)
source(s, "Sources: MyTakziah interactive prototype; reviewed JPN, KWSP and PERKESO forms in the repository research record.")
notes(
    s,
    "Here is the journey we are proposing. The family provides shared details once to MyTakziah. We would reuse them to prepare the relevant forms and a checklist of supporting documents. The family sees published conditions, known deadlines and questions requiring agency confirmation. They review each agency's information separately. In the first service, the family would still submit through official channels. Required originals, signatures and thumbprints remain. These are screens from our interactive prototype. Agency responses are simulated; it does not file claims. The QR opens that walkthrough. Our immediate goal is to test whether families arrive better prepared, with a clearer understanding of what to do next.",
    "The prototype does not file claims. Screens do not prove working AI extraction or accepted official-form exports. Once means reuse within MyTakziah, not elimination of agency formalities.",
)

# Slide 4: Innovation
s = prs.slides[3]
text_box(s, "Share the Common Evidence. Respect Every Agency Rule.", 0.78, 1.68, 11.2, 0.54, 28, INK, True)
text_box(
    s,
    "The workflow reuses what is common, then asks only for what each route still needs.",
    0.78,
    2.30,
    10.50,
    0.40,
    14.0,
    MUTED,
)
rect(s, 0.78, 2.90, 11.45, 0.92, SOFT)
text_box(s, "Shared Once", 1.02, 3.17, 1.65, 0.30, 17, TERRA, True)
for x, label in [(3.00, "Deceased Identity"), (6.05, "Applicant Identity"), (9.10, "Proof of Death")]:
    rule(s, x - 0.18, 3.10, 0.01, HAIR, 22)
    text_box(s, label, x, 3.17, 2.65, 0.30, 14, INK, True)
branches = [
    (0.78, "JPN", "Registration route, informant details and required originals."),
    (4.62, "KWSP", "Relationship evidence and payment details when the route requires them."),
    (8.46, "PERKESO", "Relationship evidence and payment details where applicable."),
]
for x, label, detail in branches:
    rule(s, x, 4.28, 3.35, SLATE, 3.0)
    text_box(s, label, x + 0.02, 4.51, 3.25, 0.35, 18, SLATE, True)
    text_box(s, detail, x + 0.02, 4.96, 3.22, 0.72, 13.5, MUTED)
rule(s, 0.78, 5.98, 11.45)
text_box(s, "AI extracts and explains. Human-maintained rules constrain. Agencies decide.", 0.78, 6.12, 11.35, 0.34, 15, INK, True)
source(s, "Sources: JPN, Ordinary Registration of Death; KWSP 9KM (AHL), V20260701; PERKESO Borang 34 and benefit checklists.")
notes(
    s,
    "The innovation is connecting a family record to the right documents and next actions across agencies. An answer from a chatbot is not a completed preparation process. MyGOV brings services together; MyDigital ID authenticates users; MyGDX supports data exchange. We propose a focused bereavement workflow that could complement them. AI would extract draft details and explain cited instructions in plain language. Rules reviewed and kept current by people would control conditions and deadlines. Uncertain information must be checked, and agencies decide eligibility. We also recognise Britain's Tell Us Once as prior art. Our focus is preparing Malaysian benefit applications, with requirements specific to these agencies. That difference still needs user testing.",
    "Official forms support recurring information categories, not identical documents. AI responsibilities and competitive differentiation are proposed design boundaries, not validated performance.",
)

# Slide 5: Feasibility
s = prs.slides[4]
text_box(s, "Test Preparation Before Government Integration", 0.78, 1.68, 11.3, 0.54, 29, INK, True)
text_box(s, "A bounded path tests usefulness and safety before any promise of automated submission.", 0.78, 2.30, 10.4, 0.38, 13.5, MUTED)
rule(s, 1.55, 3.07, 9.55, HAIR, 1.5)
stages = [
    (
        0.78,
        "Evidence Today",
        "Interactive prototype and documented agency-form research.",
        "Reality: no completed interviews, host or agency integration.",
    ),
    (
        4.48,
        "Supervised Preparation Pilot",
        "[ASSUMPTION] Test one PJ community, reviewed drafts and official channels.",
        "Gate: user, rule and privacy review; draft-format acceptance.",
    ),
    (
        8.18,
        "Authorised Integration",
        "Submission and status updates only after agency agreements and access.",
        "Gate: legal, identity, operational and technical approvals.",
    ),
]
for index, (x, label, detail, gate) in enumerate(stages):
    dot(s, x + 0.30, 2.87, 0.38, [SLATE, TERRA, HAIR][index], 4)
    text_box(s, label, x, 3.43, 3.30, 0.65, 19, INK, True)
    text_box(s, detail, x, 4.16, 3.18, 0.86, 13.5, MUTED)
    rect(s, x, 5.15, 3.25, 1.05, SOFT)
    text_box(s, gate, x + 0.18, 5.36, 2.88, 0.62, 13.0, INK, True)
source(s, "Evidence status as at 12 September 2026: prototype and form research exist; interviews, host, reviewers and integration agreements are not recorded.")
notes(
    s,
    "Today we have an interactive prototype and documented comparisons of official forms. We have no recorded user interviews or agency integration. Our proposed next step is one community pilot in Petaling Jaya, subject to securing a host. First, speak with people who have handled a death and observe where preparation breaks down. Then test a reviewed checklist and draft forms using fictional details. Before handling real documents, we need independent rule and privacy review, and confirmation that the draft format is acceptable. Only then would we offer assisted preparation through that host. This tests value before government integration. Automated submission and agency status updates would require separate agreements and technical access.",
    "The pilot, host and reviewers are assumptions or TODOs. The initial pilot does not fulfil the full Track 02 integration ambition; it is staged evidence building.",
)

# Slide 6: Impact
s = prs.slides[5]
text_box(s, "Test Whether Families Arrive Better Prepared", 0.78, 1.68, 11.5, 0.54, 28, INK, True)
for i, label in enumerate(["Less Preparation\nTime?", "Fewer Document\nCorrections?", "More Relevant\nDeadlines Met?"]):
    x = 0.78 + i * 3.82
    text_box(s, label, x, 2.72, 3.30, 1.14, 22, INK, True)
    if i < 2:
        rule(s, x + 3.52, 2.72, 0.01, HAIR, 70)
rect(s, 0.78, 4.65, 4.05, 1.28, SOFT)
text_box(s, "RM3.24b in 2024", 1.00, 4.93, 3.60, 0.36, 18, INK, True)
text_box(s, "Financial stakes in EPF death withdrawals, not MyTakziah revenue.", 1.00, 5.34, 3.58, 0.50, 12.2, MUTED)
rect(s, 5.12, 4.65, 7.10, 1.28, SOFT)
text_box(s, "[ASSUMPTION] TEST ONLY AFTER USER VALUE", 5.36, 4.91, 6.55, 0.30, 12.5, MUTED, True)
text_box(s, "Keep guidance free; test optional fixed-fee provider listings. No family-data sale, benefit commission or paid influence on guidance.", 5.36, 5.30, 6.50, 0.46, 12.2, MUTED)
source(s, "Source: EPF Integrated Annual Report 2024. Pilot outcomes and willingness to pay have not been measured.")
notes(
    s,
    "We want families to understand their next step and complete the right paperwork sooner. The pilot would measure preparation time, missing-document corrections and whether the relevant deadline was met. We have not measured those outcomes yet. EPF processed three point two four billion ringgit in death withdrawals in 2024. That shows the financial importance of this journey; it is not our revenue market. Only after demonstrating user value, our proposed funding test would keep core guidance free. Service providers could pay for optional directory listings. Families choose whether to contact them. Payment must never influence benefit guidance, and family documents must never become sales leads. Willingness to pay remains untested.",
    "RM3.24 billion is existing EPF withdrawal activity, not TAM, recoverable value or MyTakziah revenue. The directory model is an unvalidated assumption with ethical guardrails.",
)

# Slide 7: Why Us and Closing Vision
s = prs.slides[6]
text_box(s, "Help Us Test This With One Community", 0.78, 1.68, 8.05, 0.56, 30, INK, True)
points = [
    ("TEAM", "Four students: two Computer Science, two Actuarial Science."),
    ("BUILT BEFORE", "Two members built Layak, a previous hackathon grand champion."),
    ("ASK", "Community host, benefits adviser and privacy mentor."),
]
for i, (label, detail) in enumerate(points):
    y = 2.65 + i * 0.92
    text_box(s, label, 0.78, y, 1.70, 0.28, 11.2, SLATE, True)
    text_box(s, detail, 2.50, y - 0.02, 5.95, 0.62, 16, INK, True)
    if i < 2:
        rule(s, 0.78, y + 0.64, 7.75)
text_box(s, "One loss should not mean starting over at every counter.", 0.78, 5.72, 7.95, 0.82, 23, TERRA, True)
picture(s, "mark.png", 9.85, 2.66, h=1.78)
text_box(s, "MyTakziah", 9.08, 4.75, 3.30, 0.45, 25, INK, True, PP_ALIGN.CENTER)
text_box(s, "One family record. Clearer next steps.", 9.12, 5.39, 3.20, 0.50, 12, MUTED, False, PP_ALIGN.CENTER)
source(s, "Team composition and Layak experience are team-supplied records. No pilot partner is secured.")
notes(
    s,
    "Our four-person team combines two computer science and two actuarial science students. Two members also built Layak, a previous hackathon grand champion. For MyTakziah, we have a prototype and agency research. Computing supports the interface; actuarial training supports structured checks. Independent agency expertise is our next gap. Our next milestone is a reviewed preparation workflow tested with users. We ask for introductions to one community host and an experienced benefits adviser, plus a privacy mentor. Prize funding would support that work, with a budget agreed before spending. One loss should not mean starting over at every counter. Thank you.",
    "Team composition and Layak experience are team-supplied. Two members built Layak, not all four. No host, adviser or privacy mentor is secured.",
)

prs.core_properties.title = "MyTakziah Final Pitch"
prs.core_properties.subject = "URIIS Student Biz Innov Challenge 2026 Grand Final"
prs.core_properties.author = "Team Survive"
prs.save(str(OUTPUT))
print(f"saved {OUTPUT} with {len(prs.slides)} slides")
